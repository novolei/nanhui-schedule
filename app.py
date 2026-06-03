import sqlite3, os, random, json, hashlib
from datetime import datetime, timedelta
from flask import Flask, g, request, jsonify, render_template, send_from_directory

app = Flask(__name__)
DB_PATH = os.path.join(os.path.dirname(__file__), 'schedule.db')
PASSWORD_HASH = hashlib.sha256(b'admin123').hexdigest()

# ---- Database helpers ----
def get_db():
    if 'db' not in g:
        g.db = sqlite3.connect(DB_PATH)
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA journal_mode=WAL")
    return g.db

@app.teardown_appcontext
def close_db(e):
    db = g.pop('db', None)
    if db: db.close()

def init_db():
    db = get_db()
    db.executescript('''
        CREATE TABLE IF NOT EXISTS staff (
            id         INTEGER PRIMARY KEY AUTOINCREMENT,
            name       TEXT NOT NULL,
            brand      TEXT DEFAULT '',
            role       TEXT DEFAULT '营业员',
            sort_order INTEGER DEFAULT 0,
            is_active  INTEGER DEFAULT 1,
            created_at TEXT DEFAULT (datetime('now','localtime'))
        );
        CREATE TABLE IF NOT EXISTS schedules (
            id          INTEGER PRIMARY KEY AUTOINCREMENT,
            week_start  TEXT NOT NULL,
            staff_id    INTEGER REFERENCES staff(id),
            mon_shift   TEXT DEFAULT '',
            tue_shift   TEXT DEFAULT '',
            wed_shift   TEXT DEFAULT '',
            thu_shift   TEXT DEFAULT '',
            fri_shift   TEXT DEFAULT '',
            sat_shift   TEXT DEFAULT '',
            sun_shift   TEXT DEFAULT '',
            created_at  TEXT DEFAULT (datetime('now','localtime')),
            UNIQUE(week_start, staff_id)
        );
        CREATE TABLE IF NOT EXISTS published_weeks (
            week_start   TEXT PRIMARY KEY,
            published_at TEXT DEFAULT (datetime('now','localtime'))
        );
    ''')
    # seed default staff if empty
    cur = db.execute("SELECT count(*) FROM staff")
    if cur.fetchone()[0] == 0:
        names = [
            ('陈磊','销售经理','销售经理',1),
            ('刘晓庆','营业员','营业员',2),
            ('胡倩','松下','营业员',3),
            ('江凤','苏泊尔','营业员',4),
            ('陈梅芳','美的','营业员',5),
            ('郭友琴','九阳','营业员',6),
            ('刘静','石头','营业员',7),
            ('倪艺','追觅','营业员',8),
            ('杨亚男','云鲸','营业员',9),
        ]
        db.executemany("INSERT INTO staff(name,brand,role,sort_order) VALUES(?,?,?,?)", names)
        db.commit()

with app.app_context():
    init_db()

# ---- Auth ----
def check_auth(req):
    pwd = req.headers.get('X-Admin-Password', '')
    return hashlib.sha256(pwd.encode()).hexdigest() == PASSWORD_HASH

# ---- API: Staff ----
@app.route('/api/staff', methods=['GET'])
def get_staff():
    db = get_db()
    rows = db.execute("SELECT * FROM staff WHERE is_active=1 ORDER BY sort_order").fetchall()
    return jsonify([dict(r) for r in rows])

@app.route('/api/staff', methods=['POST'])
def add_staff():
    if not check_auth(request): return jsonify({'error':'密码错误'}), 401
    data = request.json
    db = get_db()
    cur = db.execute("INSERT INTO staff(name,brand,role,sort_order) VALUES(?,?,?,?)",
                     (data['name'], data.get('brand',''), data.get('role','营业员'), data.get('sort_order',99)))
    db.commit()
    return jsonify({'id': cur.lastrowid})

@app.route('/api/staff/<int:sid>', methods=['PUT'])
def update_staff(sid):
    if not check_auth(request): return jsonify({'error':'密码错误'}), 401
    data = request.json
    db = get_db()
    db.execute("UPDATE staff SET name=?, brand=?, role=?, sort_order=? WHERE id=?",
               (data['name'], data.get('brand',''), data.get('role','营业员'), data.get('sort_order',99), sid))
    db.commit()
    return jsonify({'ok':True})

@app.route('/api/staff/<int:sid>', methods=['DELETE'])
def delete_staff(sid):
    if not check_auth(request): return jsonify({'error':'密码错误'}), 401
    db = get_db()
    db.execute("UPDATE staff SET is_active=0 WHERE id=?", (sid,))
    db.commit()
    return jsonify({'ok':True})

# ---- API: Schedules ----
@app.route('/api/schedules', methods=['GET'])
def get_schedules():
    ws = request.args.get('week_start', '')
    db = get_db()
    rows = db.execute("SELECT s.*, st.name, st.brand FROM schedules s JOIN staff st ON s.staff_id=st.id WHERE s.week_start=? ORDER BY st.sort_order", (ws,)).fetchall()
    return jsonify([dict(r) for r in rows])

@app.route('/api/schedules', methods=['PUT'])
def save_schedules():
    if not check_auth(request): return jsonify({'error':'密码错误'}), 401
    data = request.json
    ws = data['week_start']
    items = data['schedules']
    db = get_db()
    for item in items:
        db.execute("""
            INSERT INTO schedules(week_start,staff_id,mon_shift,tue_shift,wed_shift,thu_shift,fri_shift,sat_shift,sun_shift)
            VALUES(?,?,?,?,?,?,?,?,?)
            ON CONFLICT(week_start,staff_id) DO UPDATE SET
                mon_shift=excluded.mon_shift, tue_shift=excluded.tue_shift,
                wed_shift=excluded.wed_shift, thu_shift=excluded.thu_shift,
                fri_shift=excluded.fri_shift, sat_shift=excluded.sat_shift,
                sun_shift=excluded.sun_shift
        """, (ws, item['staff_id'], item['mon_shift'], item['tue_shift'], item['wed_shift'],
              item['thu_shift'], item['fri_shift'], item['sat_shift'], item['sun_shift']))
    db.commit()
    return jsonify({'ok':True})

@app.route('/api/schedules/auto', methods=['POST'])
def auto_schedule():
    if not check_auth(request): return jsonify({'error':'密码错误'}), 401
    data = request.json
    ws = data['week_start']
    db = get_db()
    staff_rows = db.execute("SELECT * FROM staff WHERE is_active=1 ORDER BY sort_order").fetchall()
    staff_list = [dict(r) for r in staff_rows]
    n = len(staff_list)
    if n == 0: return jsonify({'schedules':[]})

    days = ['mon_shift','tue_shift','wed_shift','thu_shift','fri_shift','sat_shift','sun_shift']
    shifts = ['' for _ in range(7)]
    result = [shifts[:] for _ in range(n)]

    # 1) assign 休 - spread out
    idxs = list(range(n))
    random.shuffle(idxs)
    for i in range(min(n, 7)):
        result[idxs[i]][i % 7] = '休'
    for i in range(7, n):
        counts = [sum(1 for r in result if r[d]=='休') for d in range(7)]
        mn = min(counts)
        cand = [d for d,c in enumerate(counts) if c==mn]
        result[idxs[i]][random.choice(cand)] = '休'

    # 2) assign 全 - 1 per day
    used = set()
    for d in range(7):
        avail = [i for i in range(n) if i not in used and result[i][d] != '休']
        if avail:
            pick = random.choice(avail)
            result[pick][d] = '全'
            used.add(pick)
    for i in range(n):
        if i not in used:
            for d in range(7):
                if result[i][d] == '':
                    result[i][d] = '全'
                    break

    # 3) balance 早/晚 for remaining slots
    for d in range(7):
        empty = [i for i in range(n) if result[i][d] == '']
        # half early, half late
        half = len(empty) // 2
        random.shuffle(empty)
        for j, i in enumerate(empty):
            result[i][d] = '早' if j < half else '晚'

    # save to db
    db.execute("DELETE FROM schedules WHERE week_start=?", (ws,))
    for i, st in enumerate(staff_list):
        db.execute("""
            INSERT INTO schedules(week_start,staff_id,mon_shift,tue_shift,wed_shift,thu_shift,fri_shift,sat_shift,sun_shift)
            VALUES(?,?,?,?,?,?,?,?,?)
        """, (ws, st['id']) + tuple(result[i]))
    db.commit()

    # return full data
    rows = db.execute("SELECT s.*, st.name, st.brand FROM schedules s JOIN staff st ON s.staff_id=st.id WHERE s.week_start=? ORDER BY st.sort_order", (ws,)).fetchall()
    return jsonify([dict(r) for r in rows])

# ---- API: Publish ----
@app.route('/api/publish', methods=['GET'])
def get_published():
    db = get_db()
    rows = db.execute("SELECT * FROM published_weeks ORDER BY week_start DESC").fetchall()
    return jsonify([dict(r) for r in rows])

@app.route('/api/publish', methods=['POST'])
def publish():
    if not check_auth(request): return jsonify({'error':'密码错误'}), 401
    ws = request.json['week_start']
    db = get_db()
    db.execute("INSERT OR REPLACE INTO published_weeks(week_start) VALUES(?)", (ws,))
    db.commit()
    return jsonify({'ok':True})

@app.route('/api/publish/<week_start>', methods=['DELETE'])
def unpublish(week_start):
    if not check_auth(request): return jsonify({'error':'密码错误'}), 401
    db = get_db()
    db.execute("DELETE FROM published_weeks WHERE week_start=?", (week_start,))
    db.commit()
    return jsonify({'ok':True})

# ---- API: Password ----
@app.route('/api/check_password', methods=['POST'])
def check_password():
    pwd = request.json.get('password', '')
    ok = hashlib.sha256(pwd.encode()).hexdigest() == PASSWORD_HASH
    return jsonify({'ok': ok})

# ---- Share page (server-rendered with OG tags for WeChat) ----
DAYS_CN = ['周一','周二','周三','周四','周五','周六','周日']
DAY_KEYS = ['mon_shift','tue_shift','wed_shift','thu_shift','fri_shift','sat_shift','sun_shift']

@app.route('/share/<name>')
def share(name):
    db = get_db()
    # find staff
    st = db.execute("SELECT * FROM staff WHERE name=? AND is_active=1", (name,)).fetchone()
    if not st:
        return '<h2 style="text-align:center;margin-top:100px;color:#999">未找到该员工</h2>', 404

    # get latest published week
    pub = db.execute("SELECT week_start FROM published_weeks ORDER BY week_start DESC LIMIT 1").fetchone()
    if not pub:
        return '<h2 style="text-align:center;margin-top:100px;color:#999">暂无已发布的排班</h2>', 404

    ws = pub['week_start']
    row = db.execute("SELECT * FROM schedules WHERE week_start=? AND staff_id=?", (ws, st['id'])).fetchone()
    if not row:
        return '<h2 style="text-align:center;margin-top:100px;color:#999">暂无排班数据</h2>', 404

    d = datetime.strptime(ws, '%Y-%m-%d')
    week_end = d + timedelta(days=6)
    week_label = f'{d.month}月{d.day}日 - {week_end.month}月{week_end.day}日'

    days = []
    for i, label in enumerate(DAYS_CN):
        dt = d + timedelta(days=i)
        shift = row[DAY_KEYS[i]] or '-'
        days.append({'label': label, 'date': f'{dt.month}/{dt.day}', 'shift': shift})

    return render_template('share.html',
        name=name,
        week_label=week_label,
        days=days,
        base_url=request.url_root.rstrip('/'),
        request_url=request.url)


# ---- Full team share page ----
@app.route('/share/full')
def share_full():
    db = get_db()
    pub = db.execute("SELECT week_start FROM published_weeks ORDER BY week_start DESC LIMIT 1").fetchone()
    if not pub:
        return '<h2 style="text-align:center;margin-top:100px;color:#999">暂无已发布的排班</h2>', 404
    ws = pub['week_start']
    rows = db.execute("SELECT s.*, st.name, st.brand, st.role FROM schedules s JOIN staff st ON s.staff_id=st.id WHERE s.week_start=? ORDER BY st.sort_order", (ws,)).fetchall()
    d = datetime.strptime(ws, '%Y-%m-%d')
    week_end = d + timedelta(days=6)
    week_label = f'{d.month}月{d.day}日 - {week_end.month}月{week_end.day}日'
    schedule_data = []
    for r in rows:
        shifts = [r[k] or '-' for k in DAY_KEYS]
        schedule_data.append({'name': r['name'], 'brand': r['brand'] or '', 'role': r['role'] or '', 'shifts': shifts})
    dates = [(d + timedelta(days=i)) for i in range(7)]
    return render_template('share_full.html',
        week_label=week_label, schedules=schedule_data,
        days_cn=DAYS_CN, dates=dates,
        base_url=request.url_root.rstrip('/'),
        request_url=request.url)


# ---- Excel export ----
@app.route('/api/export/excel')
def export_excel():
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    ws_param = request.args.get('week_start', '')
    db = get_db()
    if not ws_param:
        pub = db.execute("SELECT week_start FROM published_weeks ORDER BY week_start DESC LIMIT 1").fetchone()
        if not pub:
            return jsonify({'error':'暂无排班数据'}), 404
        ws_param = pub['week_start']

    rows = db.execute("SELECT s.*, st.name, st.brand, st.role FROM schedules s JOIN staff st ON s.staff_id=st.id WHERE s.week_start=? ORDER BY st.sort_order", (ws_param,)).fetchall()
    staff_list = db.execute("SELECT * FROM staff WHERE is_active=1 ORDER BY sort_order").fetchall()
    d = datetime.strptime(ws_param, '%Y-%m-%d')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Sheet1'

    # Styles
    font_title = Font(name='黑体', size=20)
    font_header = Font(name='黑体', size=14)
    font_date = Font(name='黑体', size=12)
    font_data = Font(name='等线 Light', size=14)
    font_data_bold = Font(name='等线 Light', size=14, bold=True)
    align_center = Alignment(horizontal='center', vertical='center')
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    fill_gray = PatternFill(start_color='FFD9D9D9', end_color='FFD9D9D9', fill_type='solid')
    fill_white = PatternFill(start_color='FFFFFFFF', end_color='FFFFFFFF', fill_type='solid')

    # Row 1: Title
    title = f'南汇店小家电排班表（{d.month}.{d.day}-{(d+timedelta(days=6)).month}.{(d+timedelta(days=6)).day}）'
    ws.merge_cells('A1:I1')
    cell = ws['A1']
    cell.value = title
    cell.font = font_title
    cell.alignment = align_center

    # Row 2: Header (品牌, 姓名, dates)
    headers = ['品牌', '姓名']
    for i in range(7):
        dt = d + timedelta(days=i)
        headers.append(dt)  # Store as date object for serial number
    for col_idx, val in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col_idx, value=val)
        if col_idx <= 2:
            cell.font = font_header
        else:
            cell.font = font_date
            cell.number_format = 'D/M'  # Format as date like the original
        cell.alignment = align_center
        cell.border = thin_border

    # Row 3: Days of week
    day_names = ['周一','周二','周三','周四','周五','周六','周日']
    ws.cell(row=3, column=1, value='').border = thin_border
    ws.cell(row=3, column=2, value='').border = thin_border
    for i, dn in enumerate(day_names):
        cell = ws.cell(row=3, column=i+3, value=dn)
        cell.font = font_date
        cell.alignment = align_center
        cell.border = thin_border

    # Row 4+: Data
    # Determine brand/role group fills
    prev_role = None
    for idx, r in enumerate(rows):
        row_num = 4 + idx
        is_odd = idx % 2 == 1
        fill = fill_gray if is_odd else fill_white

        if r['role'] != prev_role and r['role']:
            ws.cell(row=row_num, column=1, value=r['role'])
            prev_role = r['role']

        ws.cell(row=row_num, column=1).border = thin_border
        ws.cell(row=row_num, column=1).fill = fill
        ws.cell(row=row_num, column=1).font = font_data
        ws.cell(row=row_num, column=1).alignment = align_center

        ws.cell(row=row_num, column=2, value=r['name']).font = font_data_bold
        ws.cell(row=row_num, column=2).alignment = align_center
        ws.cell(row=row_num, column=2).border = thin_border
        ws.cell(row=row_num, column=2).fill = fill

        for i in range(7):
            shift = r[DAY_KEYS[i]] or '-'
            cell = ws.cell(row=row_num, column=i+3, value=shift)
            cell.font = font_data
            cell.alignment = align_center
            cell.border = thin_border
            cell.fill = fill

    # Column widths
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 10
    for i in range(3, 10):
        ws.column_dimensions[get_column_letter(i)].width = 8

    # Save to temp file
    import tempfile
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx')
    wb.save(tmp.name)
    tmp.close()

    from flask import send_file
    return send_file(tmp.name,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=f'排班表_{ws_param}.xlsx')


# ---- Serve frontend ----
@app.route('/')
def index():
    return send_from_directory('templates', 'index.html')

@app.route('/<path:path>')
def static_files(path):
    if path.startswith('static/'):
        return send_from_directory('.', path)
    return send_from_directory('templates', 'index.html')

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    app.run(host='0.0.0.0', port=port, debug=True)
